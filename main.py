import os
import sys
import tempfile
import uuid

from dotenv import load_dotenv

load_dotenv()


# ========== 编码和 Hugging Face 设置 ==========
os.environ["PYTHONIOENCODING"] = "utf-8"
os.environ["PYTHONLEGACYWINDOWSSTDIO"] = "utf-8"
os.environ["HTTPX_DEFAULT_ENCODING"] = "utf-8"

# 使用 Hugging Face 镜像站
os.environ["HF_ENDPOINT"] = "https://hf-mirror.com"

# 延长模型下载等待时间
os.environ["HF_HUB_DOWNLOAD_TIMEOUT"] = "120"
os.environ["HF_HUB_ETAG_TIMEOUT"] = "30"

if sys.stdout.encoding and sys.stdout.encoding.lower() != "utf-8":
    sys.stdout.reconfigure(encoding="utf-8")

if sys.stderr.encoding and sys.stderr.encoding.lower() != "utf-8":
    sys.stderr.reconfigure(encoding="utf-8")


# ========== 正式导入 ==========
import streamlit as st

from langchain_community.document_loaders import PyMuPDFLoader
from langchain_text_splitters import RecursiveCharacterTextSplitter
from langchain_community.embeddings import HuggingFaceEmbeddings
from langchain_community.vectorstores import Chroma

try:
    from langchain.schema import Document
except ImportError:
    from langchain_core.documents import Document

from openai import OpenAI
import docx


# ========== 页面设置 ==========
st.set_page_config(
    page_title="RAG 文档问答",
    page_icon="📚"
)

st.title("📚 RAG 文档问答系统")
st.caption("基于 LangChain + Chroma + DeepSeek 构建")


# ========== 加载 Embedding 模型 ==========
# ========== 加载 Embedding 模型 ==========
@st.cache_resource
def load_embedding_model():
    return HuggingFaceEmbeddings(
        model_name=r"E:\t4vcollege\rag\models\bge-small-zh-v1.5",
        model_kwargs={
            "device": "cpu"
        },
        encode_kwargs={
            "normalize_embeddings": True
        }
    )


# ========== 加载 DeepSeek 客户端 ==========
@st.cache_resource
def load_client():
    api_key = os.getenv("DEEPSEEK_API_KEY")

    if not api_key:
        st.error("❌ 未检测到 DEEPSEEK_API_KEY，请检查 .env 文件。")
        st.stop()

    return OpenAI(
        api_key=api_key,
        base_url="https://api.deepseek.com/v1"
    )


# ========== 加载模型 ==========
embedding_model = load_embedding_model()
client = load_client()


# ========== 文件解析函数 ==========
def parse_file(uploaded_file):
    """
    根据文件类型解析 PDF、DOCX 和 TXT 文件。
    返回 LangChain Document 列表。
    """
    filename = uploaded_file.name
    suffix = os.path.splitext(filename)[1].lower()

    tmp_path = None

    try:
        # 将上传文件暂时保存到本地
        with tempfile.NamedTemporaryFile(
            delete=False,
            suffix=suffix
        ) as tmp:
            tmp.write(uploaded_file.getvalue())
            tmp_path = tmp.name

        docs = []

        # ---------- 解析 PDF ----------
        if suffix == ".pdf":
            loader = PyMuPDFLoader(tmp_path)
            docs = loader.load()

            # 保存真实文件名
            for doc in docs:
                doc.metadata["source"] = filename

        # ---------- 解析 Word ----------
        elif suffix == ".docx":
            docx_file = docx.Document(tmp_path)

            paragraphs = [
                paragraph.text.strip()
                for paragraph in docx_file.paragraphs
                if paragraph.text.strip()
            ]

            for i, paragraph in enumerate(paragraphs):
                docs.append(
                    Document(
                        page_content=paragraph,
                        metadata={
                            "source": filename,
                            "page": i + 1
                        }
                    )
                )

        # ---------- 解析 TXT ----------
        elif suffix == ".txt":
            with open(
                tmp_path,
                "r",
                encoding="utf-8-sig",
                errors="ignore"
            ) as f:
                full_text = f.read()

            # 按空行划分段落
            sections = [
                section.strip()
                for section in full_text.split("\n\n")
                if section.strip()
            ]

            for i, section in enumerate(sections):
                docs.append(
                    Document(
                        page_content=section,
                        metadata={
                            "source": filename,
                            "page": i + 1
                        }
                    )
                )

        return docs

    finally:
        # 无论成功还是失败，都尝试删除临时文件
        if tmp_path and os.path.exists(tmp_path):
            try:
                os.unlink(tmp_path)
            except Exception:
                pass


# ========== 初始化 Session State ==========
if "messages" not in st.session_state:
    st.session_state.messages = []

if "vectorstore" not in st.session_state:
    st.session_state.vectorstore = None

if "k_value" not in st.session_state:
    st.session_state.k_value = 3


# ========== 侧边栏 ==========
with st.sidebar:
    st.header("📂 上传文档")

    uploaded_files = st.file_uploader(
        "支持 PDF、Word、TXT",
        type=["pdf", "docx", "txt"],
        accept_multiple_files=True
    )

    if uploaded_files:
        if st.button("🔄 构建知识库"):

            with st.spinner("正在解析文件并构建知识库..."):

                all_documents = []

                # ---------- 解析所有上传文件 ----------
                for uploaded_file in uploaded_files:

                    try:
                        docs = parse_file(uploaded_file)

                        if not docs:
                            st.warning(
                                f"⚠️ 文件 {uploaded_file.name} "
                                "没有解析出有效内容。"
                            )
                            continue

                        all_documents.extend(docs)

                        st.write(
                            f"✔ 已解析：{uploaded_file.name}"
                        )

                    except Exception as e:
                        st.error(
                            f"❌ 解析文件 {uploaded_file.name} 失败：{e}"
                        )

                if not all_documents:
                    st.error(
                        "❌ 没有解析到有效文档内容，无法构建知识库。"
                    )
                    st.stop()

                # ---------- 文本切分 ----------
                splitter = RecursiveCharacterTextSplitter(
                    chunk_size=500,
                    chunk_overlap=50
                )

                chunks = splitter.split_documents(all_documents)

                if not chunks:
                    st.error(
                        "❌ 文本切分后没有得到有效文本块。"
                    )
                    st.stop()

                # ---------- 创建新的 Chroma collection ----------
                # 每次构建使用新的 collection
                collection_name = (
                    f"rag_collection_{uuid.uuid4().hex[:12]}"
                )

                vectorstore = Chroma.from_documents(
                    documents=chunks,
                    embedding=embedding_model,
                    collection_name=collection_name,
                    persist_directory="chroma_db"
                )

                # ---------- 保存本次知识库 ----------
                st.session_state.vectorstore = vectorstore

                # ---------- 重新构建后清空旧聊天记录 ----------
                st.session_state.messages = []

                st.success(
                    f"✅ 知识库构建完成，共生成 {len(chunks)} 个文本块。"
                )

                st.info(
                    "💬 旧聊天记录已清空，可以开始新的问答。"
                )

    st.divider()

    st.header("⚙️ 检索设置")

    k_value = st.slider(
        "检索块数量",
        min_value=1,
        max_value=6,
        value=3,
        help="数值越大，参考内容越多，但回答速度可能稍慢。"
    )

    st.session_state.k_value = k_value

    st.divider()

    if st.button("🗑️ 清空对话"):
        st.session_state.messages = []
        st.rerun()


# ========== 问答函数 ==========
def ask(question, vectorstore, chat_history):
    """
    先进行向量检索，再将相关内容交给 DeepSeek 生成回答。

    返回：
    answer：模型生成的回答
    sources：支持回答的文件名、页码和原文片段
    """

    k = st.session_state.get("k_value", 3)

    # ---------- 第一步：向量检索 ----------
    # 直接检索最相关的 k 个文本块
    results = vectorstore.similarity_search(
        question,
        k=k
    )

    if not results:
        return "文档中未找到相关内容。", []

    # ---------- 第二步：整理参考内容 ----------
    context_parts = [
        doc.page_content
        for doc in results
    ]

    context = "\n\n".join(context_parts)

    # ---------- 第三步：整理详细来源 ----------
    sources = []

    for doc in results:
        metadata = doc.metadata

        filename = os.path.basename(
            metadata.get("source", "未知文件")
        )

        page = metadata.get("page", 0)

        # PDF 页码从 0 开始，因此需要加 1。
        # DOCX 和 TXT 使用程序设置的虚拟页码，不再加 1。
        if (
            filename.lower().endswith(".pdf")
            and isinstance(page, int)
        ):
            display_page = page + 1
        else:
            display_page = page

        sources.append(
            {
                "filename": filename,
                "page": display_page,
                "content": doc.page_content
            }
        )

    # ---------- 第四步：去除重复来源 ----------
    unique_sources = []
    seen_sources = set()

    for source in sources:
        source_key = (
            source["filename"],
            source["page"],
            source["content"]
        )

        if source_key not in seen_sources:
            seen_sources.add(source_key)
            unique_sources.append(source)

    sources = unique_sources

    # ---------- 第五步：构造系统提示词 ----------
    system_prompt = f"""
你是一个严谨的 RAG 文档问答助手。

请严格根据下面的“参考内容”回答用户的问题。

回答要求：
1. 只能使用参考内容中的信息。
2. 不要凭空补充参考内容之外的事实。
3. 如果参考内容中没有答案，请直接回答：
   “文档中未找到相关内容。”
4. 如果问题需要结合上下文理解，可以参考之前的对话。
5. 回答应当清晰、准确，并尽量结合文档内容说明。

参考内容：
{context}
"""

    # ---------- 第六步：组织多轮对话 ----------
    messages = [
        {
            "role": "system",
            "content": system_prompt
        }
    ]

    for msg in chat_history:
        messages.append(
            {
                "role": msg["role"],
                "content": msg["content"]
            }
        )

    messages.append(
        {
            "role": "user",
            "content": question
        }
    )

    # ---------- 第七步：调用 DeepSeek ----------
    response = client.chat.completions.create(
        model="deepseek-chat",
        messages=messages
    )

    answer = response.choices[0].message.content

    return answer, sources


# ========== 显示来源函数 ==========
def show_sources(sources):
    """
    展示回答对应的文件名、页码和具体原文片段。
    """

    if not sources:
        st.write("没有可展示的来源。")
        return

    for i, source in enumerate(sources, start=1):

        st.markdown(
            f"**来源 {i}：** `{source['filename']}`"
        )

        st.write(
            f"📍 页码：第 {source['page']} 页"
        )

        st.markdown("**相关原文：**")

        st.info(source["content"])

        if i < len(sources):
            st.divider()


# ========== 聊天记录显示 ==========
for msg in st.session_state.messages:

    with st.chat_message(msg["role"]):

        st.write(msg["content"])

        if "sources" in msg:

            with st.expander("📌 来源"):
                show_sources(msg["sources"])


# ========== 聊天输入区域 ==========
if st.session_state.vectorstore is None:

    st.info(
        "👈 请先在左侧上传文件并点击“构建知识库”。"
    )

else:

    if question := st.chat_input("请输入你的问题..."):

        # ---------- 保存并显示用户问题 ----------
        st.session_state.messages.append(
            {
                "role": "user",
                "content": question
            }
        )

        with st.chat_message("user"):
            st.write(question)

        # ---------- 生成回答 ----------
        with st.chat_message("assistant"):

            with st.spinner(
                "正在检索文档并生成回答..."
            ):

                try:
                    answer, sources = ask(
                        question=question,
                        vectorstore=st.session_state.vectorstore,
                        chat_history=st.session_state.messages[:-1]
                    )

                    st.write(answer)

                    with st.expander("📌 来源"):
                        show_sources(sources)

                    # ---------- 保存助手回答和来源 ----------
                    st.session_state.messages.append(
                        {
                            "role": "assistant",
                            "content": answer,
                            "sources": sources
                        }
                    )

                except Exception as e:
                    st.error(f"❌ 生成回答时发生错误：{e}")